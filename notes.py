from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

FICHIER = Path(__file__).with_name("Etudiants.xlsx")
MATIERES = ("PC", "Français", "Anglais", "Science", "HG", "CD-ROM")
CLASSES = ("6ème", "5ème", "4ème", "3ème")


class notes:
    """Enregistre les notes d'un élève et calcule sa moyenne sur 20."""

    def __init__(self, nom, prenom, classe, valeurs):
        self.nom = nom.strip()
        self.prenom = prenom.strip()
        self.classe = classe.strip()
        if not self.nom or not self.prenom:
            raise ValueError("Le nom et le prénom sont obligatoires.")
        if self.classe not in CLASSES:
            raise ValueError("La classe doit être 6ème, 5ème, 4ème ou 3ème.")
        if len(valeurs) != len(MATIERES):
            raise ValueError("Il faut fournir une note pour chacune des six matières.")
        self.valeurs = []
        for matiere, valeur in zip(MATIERES, valeurs):
            try:
                valeur = float(valeur)
            except (TypeError, ValueError):
                raise ValueError(f"La note de {matiere} doit être un nombre.")
            if not 0 <= valeur <= 20:
                raise ValueError(f"La note de {matiere} doit être comprise entre 0 et 20.")
            self.valeurs.append(valeur)

    def enregistrer(self):
        if not FICHIER.exists():
            raise FileNotFoundError(
                f"Fichier Excel introuvable : {FICHIER}\n"
                "Placez Etudiants.xlsx dans le même dossier que les fichiers Python."
            )

        classeur = load_workbook(FICHIER)
        if "Notes" not in classeur.sheetnames:
            classeur.close()
            raise ValueError("La feuille 'Notes' est absente du fichier Excel.")
        feuille = classeur["Notes"]

        # Repère les données réelles : la validation Excel peut couvrir jusqu'à
        # la ligne 1000 même si les dernières lignes sont encore vides.
        derniere_ligne = 1
        ligne_cible = None
        for ligne in range(2, feuille.max_row + 1):
            nom_val = feuille.cell(ligne, 1).value
            prenom_val = feuille.cell(ligne, 2).value
            if nom_val is not None or prenom_val is not None:
                derniere_ligne = ligne
            nom_existant = str(nom_val or "").strip().casefold()
            prenom_existant = str(prenom_val or "").strip().casefold()
            if nom_existant == self.nom.casefold() and prenom_existant == self.prenom.casefold():
                ligne_cible = ligne
                break
        if ligne_cible is None:
            ligne_cible = derniere_ligne + 1
            for colonne in range(1, 11):
                if ligne_cible > 2:
                    feuille.cell(ligne_cible, colonne)._style = feuille.cell(ligne_cible - 1, colonne)._style

        valeurs_ligne = [self.nom, self.prenom, self.classe, *self.valeurs]
        for colonne, valeur in enumerate(valeurs_ligne, start=1):
            feuille.cell(ligne_cible, colonne, valeur)
        feuille.cell(ligne_cible, 10, f"=AVERAGE(D{ligne_cible}:I{ligne_cible})")
        feuille.cell(ligne_cible, 10).number_format = "0.00"

        # Actualise la table pour inclure les nouvelles lignes.
        for tableau in feuille.tables.values():
            if tableau.name == "TableNotes":
                tableau.ref = f"A1:J{max(feuille.max_row, ligne_cible)}"

        classeur.save(FICHIER)
        classeur.close()
        moyenne = sum(self.valeurs) / len(self.valeurs)
        print(f"Notes enregistrées pour {self.nom} {self.prenom}.")
        print(f"Moyenne des six matières : {moyenne:.2f}/20")
        return moyenne
