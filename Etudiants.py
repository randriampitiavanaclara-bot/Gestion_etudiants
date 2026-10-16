from pathlib import Path
from openpyxl import load_workbook

FICHIER = Path(__file__).with_name("Etudiants.xlsx")
CLASSES = ("6ème", "5ème", "4ème", "3ème")


class Eleve:
    """Représente un élève et permet de l'inscrire dans le fichier Excel."""

    def __init__(self, nom, prenom, classe):
        self.nom = nom.strip()
        self.prenom = prenom.strip()
        self.classe = classe.strip()
        if not self.nom or not self.prenom:
            raise ValueError("Le nom et le prénom sont obligatoires.")
        if self.classe not in CLASSES:
            raise ValueError("La classe doit être 6ème, 5ème, 4ème ou 3ème.")

    def inscrire(self):
        if not FICHIER.exists():
            raise FileNotFoundError(
                f"Fichier Excel introuvable : {FICHIER}\n"
                "Placez Etudiants.xlsx dans le même dossier que les fichiers Python."
            )

        classeur = load_workbook(FICHIER)
        if "inscription" not in classeur.sheetnames:
            classeur.close()
            raise ValueError("La feuille 'inscription' est absente du fichier Excel.")
        feuille = classeur["inscription"]

        # Recherche la dernière ligne réellement remplie (la validation Excel
        # peut étendre la feuille jusqu'à la ligne 1000 sans données).
        derniere_ligne = 1
        for ligne in range(2, feuille.max_row + 1):
            nom_val = feuille.cell(ligne, 1).value
            prenom_val = feuille.cell(ligne, 2).value
            classe_val = feuille.cell(ligne, 3).value
            if any(val is not None for val in (nom_val, prenom_val, classe_val)):
                derniere_ligne = ligne
            nom_existant = str(nom_val or "").strip().casefold()
            prenom_existant = str(prenom_val or "").strip().casefold()
            if nom_existant == self.nom.casefold() and prenom_existant == self.prenom.casefold():
                classeur.close()
                print("Cet élève est déjà inscrit dans le fichier Excel.")
                return False

        nouvelle_ligne = derniere_ligne + 1
        valeurs = [self.nom, self.prenom, self.classe]
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(nouvelle_ligne, colonne, valeur)
            if nouvelle_ligne > 2:
                cellule._style = feuille.cell(nouvelle_ligne - 1, colonne)._style

        # Ajuste la plage du tableau Excel pour inclure la nouvelle ligne.
        for tableau in feuille.tables.values():
            if tableau.name == "TableInscription":
                tableau.ref = f"A1:C{nouvelle_ligne}"

        classeur.save(FICHIER)
        classeur.close()
        print(f"Élève inscrit avec succès : {self.nom} {self.prenom} ({self.classe}).")
        return True
